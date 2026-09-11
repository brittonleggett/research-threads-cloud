import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

d = docx.Document()


def h1(text):
    d.add_heading(text, level=1)


def h2(text):
    d.add_heading(text, level=2)


def h3(text):
    d.add_heading(text, level=3)


def p(text, bold=False, italic=False):
    para = d.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    return para


def bullets(items):
    for item in items:
        d.add_paragraph(item, style='List Bullet')


def numbered(items):
    for item in items:
        d.add_paragraph(item, style='List Number')


# ---------------------------------------------------------------- Title page
title = d.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("MCNEESE STATE UNIVERSITY")
r.bold = True
r.font.size = Pt(16)

sub = d.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Human Subjects Institutional Review Board\nDATA COLLECTION INSTRUMENTS")
r.bold = True
r.font.size = Pt(13)

sub2 = d.add_paragraph()
sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub2.add_run("Consumer Responses to Tariff-Related Price Increase Messaging")
r.italic = True
r.font.size = Pt(12)

d.add_paragraph()
p("Principal Investigator: Britton R. Leggett")
p("Department: College of Business, McNeese State University")
p("Campus Address: Burton Business Center, 4205 Ryan Street, Lake Charles, LA 70605")
p("Email: bleggett1@mcneese.edu")
p("Phone: 337-475-5578 (ext. 5578)")
p("Date Prepared: 2026-09-08")
p("Attachment to: HSIRB Application Package -- Consumer Responses to Tariff-Related Price "
  "Increase Messaging")

d.add_paragraph()
p("This document contains the full data collection instruments for all three human-subjects "
  "phases of this research program: the vignette Pretest, Study 2 (the main vignette "
  "experiment), and Study 3 (the SEM validation survey). Study 1 uses only publicly available "
  "corporate documents and involves no human subjects; it is not included here. All three "
  "instruments below will be administered online via Qualtrics to Prolific panel respondents. "
  "Informed consent language is provided in the separate document "
  "Informed_Consent_Tariffs.docx and is referenced, not reproduced, below.")

d.add_page_break()

# ---------------------------------------------------------------- Shared vignette text
VIGNETTE_INTRO = (
    "Imagine you are a regular customer of Meridian Home, a national retailer where you often "
    "shop. You receive the following email from the company:\n"
    "Subject: An Update on Pricing at Meridian Home"
)

VIGNETTES = [
    ("Cell 1 -- Tariff-explicit x Full pass-through",
     "As you may have seen in the news, new tariffs on imported goods have significantly "
     "increased the cost of many of the products we sell. Because of these tariff-related "
     "costs, we are raising prices on select items starting next month. These increases "
     "reflect the full amount of the added tariff costs we are now paying to bring these "
     "products to you."),
    ("Cell 2 -- Tariff-explicit x Shared-burden",
     "As you may have seen in the news, new tariffs on imported goods have significantly "
     "increased the cost of many of the products we sell. Because of these tariff-related "
     "costs, we are raising prices on select items starting next month. We are absorbing a "
     "significant portion of these added tariff costs ourselves, and only passing along part "
     "of the increase to our customers."),
    ("Cell 3 -- General-cost-explicit x Full pass-through",
     "Due to rising costs across our supply chain, the cost of many of the products we sell "
     "has significantly increased. Because of these higher costs, we are raising prices on "
     "select items starting next month. These increases reflect the full amount of the added "
     "costs we are now paying to bring these products to you."),
    ("Cell 4 -- General-cost-explicit x Shared-burden",
     "Due to rising costs across our supply chain, the cost of many of the products we sell "
     "has significantly increased. Because of these higher costs, we are raising prices on "
     "select items starting next month. We are absorbing a significant portion of these "
     "added costs ourselves, and only passing along part of the increase to our customers."),
    ("Cell 5 -- Silent x Full pass-through",
     "We want to keep you informed about some upcoming changes here at Meridian Home. "
     "Starting next month, prices on select items across our stores will be increasing. We "
     "are passing along the full amount of this price increase to our customers. We "
     "appreciate your continued business and look forward to serving you as always."),
    ("Cell 6 -- Silent x Shared-burden",
     "We want to keep you informed about some upcoming changes here at Meridian Home. "
     "Starting next month, prices on select items across our stores will be increasing. We "
     "are absorbing a significant portion of this price increase ourselves, and passing along "
     "only part of it to our customers. We appreciate your continued business and look "
     "forward to serving you as always."),
]


def add_vignette_block():
    p("Shared framing shown to every respondent before their randomly assigned single cell "
      "(each respondent sees only one of the six versions below):", italic=True)
    p(VIGNETTE_INTRO)
    for title_txt, body_txt in VIGNETTES:
        h3(title_txt)
        p(body_txt)


# ================================================================== Instrument 1: Pretest
h1("Instrument 1: Stimulus Pretest (Vignette Validation)")
p("Purpose: to confirm the six vignette manipulations produce the intended differences in "
  "attribution and cost-response perceptions before they are used in Study 2, and to rule out "
  "length, likability, and realism confounds across cells. Target N = 150-180 (~25-30 per "
  "cell), Prolific, estimated completion time ~6-7 minutes.", italic=True)

h2("1. Consent")
p("Electronic informed consent per Informed_Consent_Tariffs.docx. \"I agree to participate\" / "
  "\"I do not agree to participate\"; non-consenting respondents redirected out immediately.")

h2("2. Screener")
bullets([
    "Age (open numeric entry) -- terminate if under 18.",
    "\"Do you currently reside in the United States?\" -- terminate if No.",
    "An attention/English-fluency check item embedded naturally in the screener.",
])

h2("3. Random assignment")
p("Qualtrics randomizer assigns each respondent to exactly one of the six cells below with "
  "equal probability.")

h2("4. Vignette exposure")
add_vignette_block()

h2("5. Manipulation checks")
bullets([
    "\"According to the message, what did the company say was the reason for the price "
    "increase?\" -- forced choice: tariffs / general rising costs / no reason given / don't "
    "recall.",
    "\"How clearly did the message explain the reason for the price increase?\" -- 7-point "
    "scale, 1 = not at all clearly to 7 = very clearly.",
    "\"According to the message, is the company passing along the full cost increase, or "
    "absorbing part of it?\" -- forced choice: full amount / partial, company absorbing some "
    "/ unclear.",
])

h2("6. Confound checks")
p("These items should show no significant difference across the six cells; a difference here "
  "would indicate the vignettes are confounded on something other than the intended "
  "manipulations.", italic=True)
bullets([
    "Perceived message length/complexity: \"How long or complex did this message seem to "
    "you?\" -- 7-point scale, 1 = very short/simple to 7 = very long/complex.",
    "Perceived company likability/credibility: 2-3 items adapted from the Source Credibility "
    "Model (Hovland & Weiss, 1951), e.g. \"How likable does this company seem?\" and \"How "
    "credible does this message seem?\" -- 7-point scales. Exact final wording to be "
    "confirmed by the PI before fielding; presented here as representative item content.",
    "Perceived realism: \"This message reads like something a real company would actually "
    "send.\" -- 7-point scale, 1 = strongly disagree to 7 = strongly agree.",
])

h2("7. Debriefing")
p("Thank you for completing this study.", bold=True)
p("This study is testing several different versions of a business email about a price "
  "increase, to see which version reads most clearly and realistically to people like you. "
  "\"Meridian Home\" is not a real company; the message was written for research purposes "
  "only. Your responses will help refine the materials used in a later phase of this "
  "research.")
p("Researcher contact: Britton R. Leggett, bleggett1@mcneese.edu, 337-475-5578")
p("IRB contact (questions about your rights as a research participant): Dr. Eddie Lyons, IRB "
  "Chair, irbchair@mcneese.edu, 337-475-4077")

d.add_page_break()

# ================================================================== Instrument 2: Study 2
h1("Instrument 2: Study 2 -- Vignette Experiment")
p("3 (Attribution Frame: tariff-explicit / general-cost-explicit / silent) x 2 (Cost-Response: "
  "full pass-through / shared-burden absorption) between-subjects factorial. Target N = "
  "360-600, Prolific, estimated completion time ~12 minutes.", italic=True)

h2("1. Consent")
p("Electronic informed consent per Informed_Consent_Tariffs.docx (~12-minute time estimate). "
  "\"I agree to participate\" / \"I do not agree to participate\"; non-consenting respondents "
  "redirected out immediately.")

h2("2. Screener")
bullets([
    "Age (open numeric entry) -- terminate if under 18.",
    "\"Do you currently reside in the United States?\" -- terminate if No.",
    "An attention/English-fluency check item embedded naturally in the screener.",
])

h2("3. Random assignment")
p("Qualtrics randomizer assigns each respondent to exactly one of the six cells below with "
  "equal probability. A brief forced pause / minimum time-on-page (e.g., 10 seconds) is set "
  "on the vignette-exposure page to discourage skimming.")

h2("4. Vignette exposure")
add_vignette_block()

h2("5. Manipulation checks")
p("Retained in the main study for data-quality screening, not only in the pretest.",
  italic=True)
bullets([
    "\"According to the message, what did the company say was the reason for the price "
    "increase?\" -- forced choice: tariffs / general rising costs / no reason given / don't "
    "recall.",
    "\"How clearly did the message explain the reason for the price increase?\" -- 7-point "
    "scale, 1 = not at all clearly to 7 = very clearly.",
    "\"According to the message, is the company passing along the full cost increase, or "
    "absorbing part of it?\" -- forced choice: full amount / partial, company absorbing some "
    "/ unclear.",
])

h2("6. Perceived Price Fairness -- Campbell (1999)")
p("Two items, averaged (item 2 reverse-scored before averaging):")
numbered([
    "\"This price is:\" -- 7-point bipolar scale, 1 = very fair to 7 = very unfair.",
    "\"This price is not fair.\" -- 7-point scale, 1 = strongly agree to 7 = strongly "
    "disagree (reverse-scored).",
])
p("Reported reliability in the original source: r = .84, p < .0001.", italic=True)

h2("7. Perceived Opportunism -- Campbell (1999, 2007)")
numbered([
    "\"The motive behind this price increase was:\" -- 7-point scale, 1 = bad to 7 = good.",
    "\"The intent in this situation was to take advantage of you (the customer).\" -- 7-point "
    "scale, 1 = agree to 7 = disagree.",
])
p("Note: item 2's polarity (agree = low end) is reversed relative to the other scales in this "
  "instrument. Anchors will be flipped for respondent-facing consistency, or the item will be "
  "reverse-scored explicitly during analysis.", italic=True)

h2("8. Trust in the Company -- Chaudhuri & Holbrook (2001)")
p("Four items, 7-point Likert (1 = strongly disagree to 7 = strongly agree):")
bullets([
    "I trust this company.",
    "I rely on this company.",
    "This is an honest company.",
    "This company is safe.",
])

h2("9. Purchase Intention -- Dodds, Monroe & Grewal (1991)")
p("Five items, adapted from \"product\" wording in the original (all 7-point scales):")
numbered([
    "The likelihood of purchasing this product is: (very high to very low)",
    "If I were going to buy this product, I would consider buying this product at the price "
    "shown. (strongly agree to strongly disagree)",
    "At the price shown, I would consider buying the product. (strongly agree to strongly "
    "disagree)",
    "The probability that I would consider buying the product is: (very high to very low)",
    "My willingness to buy the product is: (very high to very low)",
])
p("Reported reliability in the original source: coefficient alpha .96-.97. Note: items 2-3 "
  "reference \"the price shown,\" which overlaps in content with the Fairness scale above; "
  "this is a deliberate design choice, checked empirically via a discriminant-validity "
  "analysis once data are collected.", italic=True)

h2("10. Word-of-Mouth Intention -- Maxham & Netemeyer (2002)")
p("Three items, genericized from the original's banking-services wording:")
numbered([
    "How likely are you to spread positive word-of-mouth about [company]? (7-point, very "
    "unlikely to very likely)",
    "I would recommend [company]'s products to my friends. (strongly disagree to strongly "
    "agree)",
    "If my friends were looking for a product like this, I would tell them to try [company]. "
    "(strongly disagree to strongly agree)",
])

h2("11. Demographics")
bullets([
    "Age -- open numeric entry.",
    "Gender -- Male / Female / Non-binary / Prefer to self-describe / Prefer not to say.",
    "Household income -- <$25k / $25k-49,999 / $50k-74,999 / $75k-99,999 / $100k-149,999 / "
    "$150k+ / Prefer not to say.",
    "Education -- Less than high school / High school or GED / Some college, no degree / "
    "Associate's / Bachelor's / Graduate or professional degree.",
    "Shopping frequency at retailers like the one described -- Never / Rarely (a few times a "
    "year) / Occasionally (monthly) / Regularly (a few times a month) / Frequently (weekly or "
    "more).",
    "Political ideology -- single 7-point item, 1 = Very liberal to 7 = Very conservative.",
])

h2("12. Debriefing")
p("Thank you for completing this study.", bold=True)
p("We want to tell you a bit more about what this research is really about. This study "
  "examines how the way a company explains a price increase -- specifically, whether it "
  "names tariffs, cites general rising costs, or gives no explanation at all, and whether it "
  "says it's passing along the full cost or absorbing part of it -- affects how fair, "
  "trustworthy, and worth buying from consumers find that company. The message you read from "
  "\"Meridian Home\" was one of several different versions used in this study; Meridian Home "
  "is not a real company, and the email was written for research purposes only.")
p("We didn't tell you this in advance because knowing the specific comparison being tested "
  "can change how people respond, which would make the results less accurate. This is a "
  "standard and accepted research practice, not deception in the sense of anything false "
  "being told to you -- everything you were told about the study (that you'd read a business "
  "message and answer questions about it) was true.")
p("Your responses are completely anonymous and cannot be linked back to you. If anything "
  "about this study concerns you, or if you'd like to withdraw your data even after "
  "completing the survey, you can contact the research team below and reference the date and "
  "approximate time you completed the study.")
p("Researcher contact: Britton R. Leggett, bleggett1@mcneese.edu, 337-475-5578")
p("IRB contact (questions about your rights as a research participant): Dr. Eddie Lyons, IRB "
  "Chair, irbchair@mcneese.edu, 337-475-4077")

d.add_page_break()

# ================================================================== Instrument 3: Study 3
h1("Instrument 3: Study 3 -- SEM Validation Survey (No Manipulation)")
p("Correlational survey, no experimental manipulation; participants report on a real recent "
  "experience to preserve natural variance in the constructs for model validation. Target N = "
  "300-400, Prolific, estimated completion time ~15 minutes.", italic=True)

h2("1. Consent")
p("Electronic informed consent per Informed_Consent_Tariffs.docx (~15-minute time estimate). "
  "\"I agree to participate\" / \"I do not agree to participate\"; non-consenting respondents "
  "redirected out immediately.")

h2("2. Screener")
bullets([
    "Age (open numeric entry) -- terminate if under 18.",
    "\"Do you currently reside in the United States?\" -- terminate if No.",
    "An attention/English-fluency check item embedded naturally in the screener.",
])

h2("3. Recall prompt")
p("\"Think of a time in the last 12 months when a retailer you shop with raised prices on "
  "something you buy. Please briefly describe what happened.\" (open text) -- If a "
  "participant indicates they cannot recall any such instance, they are branched to a brief "
  "generic-attitude version of the items below instead of being forced to fabricate a "
  "memory.")

h2("4. Perceived Price Fairness -- Campbell (1999)")
p("Same two items as Study 2, reworded to reference \"this price increase\"/\"the retailer\" "
  "instead of the vignette scenario.")

h2("5. Perceived Opportunism -- Campbell (1999, 2007)")
p("Same two items as Study 2, reworded to reference the recalled experience.")

h2("6. Trust in the Company -- Chaudhuri & Holbrook (2001)")
p("Same four items as Study 2, reworded to reference \"this retailer.\"")

h2("7. Purchase Intention -- Dodds, Monroe & Grewal (1991)")
p("Same five items as Study 2, reworded to reference the recalled retailer/product.")

h2("8. Word-of-Mouth Intention -- Maxham & Netemeyer (2002)")
p("Same three items as Study 2, reworded to reference the recalled retailer.")

h2("9. Demographics")
p("Same battery as Study 2 (age, gender, household income, education, shopping frequency, "
  "political ideology), plus one additional covariate:")
bullets([
    "Recency and category of the recalled price increase (e.g., how long ago it occurred and "
    "what type of product/retailer it involved) -- used to rule out category-specific "
    "confounds in PLS-SEM robustness checks.",
])

h2("10. Debriefing")
p("Thank you for completing this study.", bold=True)
p("This study examines how consumers' perceptions of fairness and trust after a real price "
  "increase relate to their intentions to keep buying from and recommending that retailer. "
  "There is no hidden manipulation in this survey -- you were asked to describe and answer "
  "questions about an actual experience of your own choosing.")
p("Your responses are completely anonymous and cannot be linked back to you. If anything "
  "about this study concerns you, you can contact the research team below.")
p("Researcher contact: Britton R. Leggett, bleggett1@mcneese.edu, 337-475-5578")
p("IRB contact (questions about your rights as a research participant): Dr. Eddie Lyons, IRB "
  "Chair, irbchair@mcneese.edu, 337-475-4077")

d.add_page_break()

# ================================================================== Citations
h1("Scale Citations")
bullets([
    "Campbell, M. C. (1999). Perceptions of price unfairness: Antecedents and consequences. "
    "Journal of Marketing Research, 36(2), 187-199. [Perceived Price Fairness, Perceived "
    "Opportunism item 1 -- item wording verified against original source]",
    "Campbell, M. C. (2007). \"Says who?!\" How the source of price information and "
    "affect influence perceived price (un)fairness. Journal of Marketing Research. "
    "[Perceived Opportunism item 2 -- operationalization of the 1999 inferred-motive "
    "construct; best available sourcing, see project notes for verification detail]",
    "Chaudhuri, A., & Holbrook, M. B. (2001). The chain of effects from brand trust and brand "
    "affect to brand performance: The role of brand loyalty. Journal of Marketing, 65(2), "
    "81-93. [Trust in the Company -- verified]",
    "Dodds, W. B., Monroe, K. B., & Grewal, D. (1991). Effects of price, brand, and store "
    "information on buyers' product evaluations. Journal of Marketing Research, 28(3), "
    "307-319. [Purchase Intention -- verified verbatim against original appendix]",
    "Maxham, J. G., & Netemeyer, R. G. (2002). Journal of Marketing, 66(4), 57-71. "
    "[Word-of-Mouth Intention -- verified verbatim against original appendix]",
    "Hovland, C. I., & Weiss, W. (1951). The influence of source credibility on "
    "communication effectiveness. Public Opinion Quarterly, 15(4), 635-650. [Pretest-only "
    "confound-check items -- illustrative wording, not the focal measures of this study]",
])

d.save("../Tariff_Data_Collection_Instruments_2026-09-08.docx")
print("Saved.")
